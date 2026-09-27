"""
Specialized DOCX Builder for POLAR_GRID_SIH_NOTES.docx
Transforms docs/POLAR_GRID_SIH_NOTES.md into a publication-grade, professionally styled Word document.
Features:
- Title Page / Executive Header
- Custom colored Callout Boxes (IMPORTANT, JUDGE QUESTION, BEST ANSWER, IN SIMPLE WORDS)
- Clean Navy Tables with alternating zebra striping
- Monospace Code blocks
- Professional typography (Calibri / Consolas, consistent margins)
"""

import os
import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DOCS_DIR = os.path.join(PROJECT_ROOT, "docs")
MD_PATH = os.path.join(DOCS_DIR, "POLAR_GRID_SIH_NOTES.md")
DOCX_PATH = os.path.join(DOCS_DIR, "POLAR_GRID_SIH_NOTES.docx")

# Color Palette
HEX_NAVY = "0F172A"       # Primary Dark Navy
HEX_OCEAN = "0369A1"      # Primary Blue Accent
HEX_TEAL = "0F766E"       # Teal Subheading
HEX_AMBER_BG = "FEF3C7"   # Important Box Fill
HEX_AMBER_BORDER = "D97706"
HEX_INDIGO_BG = "EEF2FF"  # Judge Question Fill
HEX_INDIGO_BORDER = "4F46E5"
HEX_EMERALD_BG = "ECFDF5" # Best Answer Fill
HEX_EMERALD_BORDER = "059669"
HEX_SKY_BG = "F0F9FF"     # In Simple Words Fill
HEX_SKY_BORDER = "0284C7"
HEX_GRAY_BG = "F8FAFC"    # Default Box / Zebra Fill
HEX_GRAY_BORDER = "CBD5E1"
HEX_CODE_BG = "F1F5F9"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for side, border_style in borders.items():
        if border_style:
            val, sz, color = border_style
            elem = parse_xml(f'<w:{side} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(elem)
        else:
            elem = parse_xml(f'<w:{side} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(elem)
            
    tcPr.append(tcBorders)

def add_styled_paragraph(doc, text="", style='Normal', space_before=0, space_after=4, line_spacing=1.15):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    return p

def parse_inline_markdown(paragraph, text, base_font_size=10.5, is_italic=False, is_bold=False, font_color=None):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))', text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = 'Calibri'
        run.font.size = Pt(base_font_size)
        if font_color:
            run.font.color.rgb = font_color

        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run.text = token[2:-2]
            run.bold = True
            run.italic = is_italic
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run.text = token[1:-1]
            run.italic = True
            run.bold = is_bold
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run.text = token[1:-1]
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 1)
            run.font.color.rgb = RGBColor(190, 24, 93) # Pink/Magenta code font
        elif token.startswith('[') and ']' in token and '(' in token and token.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                label, url = m.groups()
                run.text = label
                run.font.color.rgb = RGBColor(2, 132, 199)
                run.underline = True
            else:
                run.text = token
        else:
            run.text = token
            run.bold = is_bold
            run.italic = is_italic

def add_callout_box(doc, title, text_lines, box_type='note'):
    """
    Creates a styled callout box (table with left accent border and colored background).
    """
    if box_type == 'important':
        bg_hex = HEX_AMBER_BG
        border_hex = HEX_AMBER_BORDER
        title_color = RGBColor(180, 83, 9)
    elif box_type == 'judge_question':
        bg_hex = HEX_INDIGO_BG
        border_hex = HEX_INDIGO_BORDER
        title_color = RGBColor(67, 56, 202)
    elif box_type == 'best_answer':
        bg_hex = HEX_EMERALD_BG
        border_hex = HEX_EMERALD_BORDER
        title_color = RGBColor(4, 120, 87)
    elif box_type == 'simple_words':
        bg_hex = HEX_SKY_BG
        border_hex = HEX_SKY_BORDER
        title_color = RGBColor(3, 105, 161)
    else:
        bg_hex = HEX_GRAY_BG
        border_hex = HEX_GRAY_BORDER
        title_color = RGBColor(51, 65, 85)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(7.0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    # Left thick border, no other borders
    set_cell_borders(cell, left=('single', '36', border_hex))

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.font.name = 'Calibri'
        run_title.font.size = Pt(11)
        run_title.bold = True
        run_title.font.color.rgb = title_color

    for idx, tl in enumerate(text_lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
        parse_inline_markdown(p, tl, base_font_size=10.5, font_color=RGBColor(30, 41, 59))

    # Add small spacer after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def build_docx():
    print(f"Reading {MD_PATH}...")
    with open(MD_PATH, "r", encoding="utf-8") as f:
        content = f.read()

    doc = Document()

    # Set page margins to 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Title Page / Cover Section
    title_p = add_styled_paragraph(doc, space_before=24, space_after=6, line_spacing=1.1)
    run_t = title_p.add_run("POLAR GRID")
    run_t.font.name = 'Calibri'
    run_t.font.size = Pt(32)
    run_t.bold = True
    run_t.font.color.rgb = RGBColor(15, 23, 42) # Slate 900

    sub_p = add_styled_paragraph(doc, space_before=2, space_after=12)
    run_s = sub_p.add_run("AI-Assisted Renewable Energy Forecasting & Microgrid Dispatch Optimization for Antarctic Research Stations")
    run_s.font.name = 'Calibri'
    run_s.font.size = Pt(16)
    run_s.font.color.rgb = RGBColor(3, 105, 161) # Ocean Blue

    loc_p = add_styled_paragraph(doc, space_before=0, space_after=18)
    run_l = loc_p.add_run("Target Case Study: Mawson Station, Mac. Robertson Land, Antarctica (-67.6027° S, 62.8738° E, Elevation 16m MSL)")
    run_l.font.name = 'Calibri'
    run_l.font.size = Pt(12)
    run_l.italic = True
    run_l.font.color.rgb = RGBColor(71, 85, 105)

    # Metadata Banner Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Event / Purpose", "Smart India Hackathon (SIH) — Student Master Defense & Notes Guide"),
        ("Language Style", "Conversational Speakable Hinglish (Technical terms in English with simple explanations)"),
        ("Source of Truth", "Actual Polar Grid Implementation (src/, backend/, frontend/, config/, tests/)"),
        ("Verified Models", "HistGradientBoostingRegressor (ML Load) + Open-Meteo ECMWF IFS (Live Weather) + SciPy HiGHS (LP Dispatch)"),
        ("Document Status", "Production-Ready SIH Defense Guide (26 Comprehensive Parts & 66 Detailed Q&As)")
    ]
    for r_idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[r_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(5.0)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, 80, 80, 120, 120)
        set_cell_margins(c1, 80, 80, 120, 120)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(k)
        r0.font.name = 'Calibri'
        r0.font.size = Pt(10)
        r0.bold = True
        r0.font.color.rgb = RGBColor(15, 23, 42)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(v)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

    # Parse body
    lines = content.split('\n')
    in_code_block = False
    code_block_lines = []
    in_table = False
    table_lines = []
    in_callout = False
    callout_type = 'note'
    callout_title = ''
    callout_lines = []

    def flush_code():
        nonlocal in_code_block, code_block_lines
        if not code_block_lines:
            in_code_block = False
            return
        code_text = '\n'.join(code_block_lines)
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(7.0)
        set_cell_background(cell, HEX_CODE_BG)
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        set_cell_borders(cell, left=('single', '12', 'CBD5E1'), top=('single', '6', 'E2E8F0'), right=('single', '6', 'E2E8F0'), bottom=('single', '6', 'E2E8F0'))
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.0)
        run.font.color.rgb = RGBColor(30, 41, 59)
        
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(3)
        code_block_lines = []
        in_code_block = False

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return
        parsed_rows = []
        for tl in table_lines:
            tl_strip = tl.strip()
            if not tl_strip:
                continue
            if tl_strip.startswith('|'):
                tl_strip = tl_strip[1:]
            if tl_strip.endswith('|'):
                tl_strip = tl_strip[:-1]
            cols = [c.strip() for c in tl_strip.split('|')]
            if all(re.match(r'^:?-+:?$', c) for c in cols if c):
                continue
            parsed_rows.append(cols)

        if parsed_rows:
            num_cols = max(len(r) for r in parsed_rows)
            for r in parsed_rows:
                while len(r) < num_cols:
                    r.append("")

            table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True

            for row_idx, row_data in enumerate(parsed_rows):
                is_header = (row_idx == 0)
                tr = table.rows[row_idx]
                for col_idx, cell_val in enumerate(row_data):
                    cell = tr.cells[col_idx]
                    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                    if is_header:
                        set_cell_background(cell, HEX_NAVY)
                    else:
                        if row_idx % 2 == 1:
                            set_cell_background(cell, "F8FAFC")
                        else:
                            set_cell_background(cell, "FFFFFF")
                    
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    clean_val = cell_val.replace('\\|', '|').strip()
                    if is_header:
                        run = p.add_run(clean_val)
                        run.bold = True
                        run.font.name = 'Calibri'
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        parse_inline_markdown(p, clean_val, base_font_size=9.5)

            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(4)

        table_lines = []
        in_table = False

    def flush_callout():
        nonlocal in_callout, callout_type, callout_title, callout_lines
        if not callout_lines:
            in_callout = False
            return
        add_callout_box(doc, callout_title, callout_lines, box_type=callout_type)
        callout_lines = []
        callout_title = ''
        in_callout = False

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle Code Block
        if stripped.startswith('```'):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                if in_callout:
                    flush_callout()
                in_code_block = True
                code_block_lines = []
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Handle Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            if in_callout:
                flush_callout()
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            flush_table()

        # Handle Blockquotes / Callout Boxes
        if stripped.startswith('>'):
            raw_content = stripped[1:].strip()
            
            # Check if this is the start of a new callout
            if not in_callout:
                in_callout = True
                callout_lines = []
                # Determine type
                if "IMPORTANT" in raw_content.upper():
                    callout_type = 'important'
                    callout_title = '⚡ IMPORTANT NOTE / CRITICAL DISTINCTION'
                elif "JUDGE QUESTION" in raw_content.upper() or "QUESTION" in raw_content.upper():
                    callout_type = 'judge_question'
                    callout_title = '🎤 JUDGE QUESTION'
                elif "BEST ANSWER" in raw_content.upper() or "JUDGE-FRIENDLY" in raw_content.upper():
                    callout_type = 'best_answer'
                    callout_title = '✅ BEST / READY-TO-SPEAK ANSWER'
                elif "SIMPLE WORDS" in raw_content.upper() or "EXPLAIN" in raw_content.upper():
                    callout_type = 'simple_words'
                    callout_title = '💡 IN SIMPLE WORDS'
                else:
                    callout_type = 'note'
                    callout_title = '📌 KEY TAKEAWAY'
            
            # Clean content if it starts with the header
            cleaned_line = raw_content
            for prefix in ['**IMPORTANT**:', '**IMPORTANT**', '[!IMPORTANT]', '**JUDGE QUESTION**:', '**JUDGE QUESTION**', '**BEST ANSWER**:', '**BEST ANSWER**', '**IN SIMPLE WORDS**:', '**IN SIMPLE WORDS**']:
                if cleaned_line.startswith(prefix):
                    cleaned_line = cleaned_line[len(prefix):].strip()
                    break
            
            if cleaned_line:
                callout_lines.append(cleaned_line)
            i += 1
            continue
        elif in_callout:
            flush_callout()

        # Empty line
        if not stripped:
            i += 1
            continue

        # Horizontal Rule
        if re.match(r'^(\-{3,}|\*{3,}|_{3,})$', stripped):
            p = add_styled_paragraph(doc, space_before=6, space_after=6)
            run = p.add_run("―" * 55)
            run.font.color.rgb = RGBColor(203, 213, 225)
            run.font.size = Pt(9)
            i += 1
            continue

        # Heading 1
        if stripped.startswith('# '):
            h_text = stripped[2:].strip()
            # If it's a major PART heading, give it a prominent style
            p = add_styled_paragraph(doc, space_before=16, space_after=6)
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(18)
            run.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42) # Navy
            i += 1
            continue

        # Heading 2
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = add_styled_paragraph(doc, space_before=12, space_after=4)
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(14)
            run.bold = True
            run.font.color.rgb = RGBColor(3, 105, 161) # Ocean blue
            i += 1
            continue

        # Heading 3
        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = add_styled_paragraph(doc, space_before=10, space_after=3)
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(15, 118, 110) # Teal
            i += 1
            continue

        # Heading 4
        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = add_styled_paragraph(doc, space_before=8, space_after=2)
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Bullet Lists
        bullet_match = re.match(r'^(\s*)([-*+])\s+(.*)$', line)
        if bullet_match:
            indent_spaces, _, item_text = bullet_match.groups()
            indent_level = len(indent_spaces) // 2
            p = add_styled_paragraph(doc, space_before=1, space_after=2)
            p.paragraph_format.left_indent = Inches(0.25 + 0.2 * indent_level)
            bullet_char = "• " if indent_level == 0 else "⁃ "
            run_b = p.add_run(bullet_char)
            run_b.font.name = 'Calibri'
            run_b.font.size = Pt(11)
            run_b.font.color.rgb = RGBColor(3, 105, 161)
            parse_inline_markdown(p, item_text, base_font_size=10.5)
            i += 1
            continue

        # Numbered Lists
        num_match = re.match(r'^(\s*)(\d+[\.\)])\s+(.*)$', line)
        if num_match:
            indent_spaces, num_prefix, item_text = num_match.groups()
            indent_level = len(indent_spaces) // 2
            p = add_styled_paragraph(doc, space_before=1, space_after=2)
            p.paragraph_format.left_indent = Inches(0.25 + 0.2 * indent_level)
            run_n = p.add_run(num_prefix + " ")
            run_n.font.name = 'Calibri'
            run_n.font.size = Pt(10.5)
            run_n.bold = True
            run_n.font.color.rgb = RGBColor(3, 105, 161)
            parse_inline_markdown(p, item_text, base_font_size=10.5)
            i += 1
            continue

        # Normal Paragraph
        p = add_styled_paragraph(doc, space_before=2, space_after=4)
        parse_inline_markdown(p, stripped, base_font_size=10.5)
        i += 1

    if in_code_block:
        flush_code()
    if in_table:
        flush_table()
    if in_callout:
        flush_callout()

    print(f"Saving {DOCX_PATH}...")
    doc.save(DOCX_PATH)
    print(f"Successfully generated {DOCX_PATH}!")

if __name__ == "__main__":
    build_docx()

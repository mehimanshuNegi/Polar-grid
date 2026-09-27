import os
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_styled_paragraph(doc, text="", style='Normal', space_before=0, space_after=6, line_spacing=1.15):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    return p

def parse_inline_markdown(paragraph, text, base_font_size=11, is_italic=False, is_bold=False, font_color=None):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))', text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = 'Calibri'
        run.font.size = Pt(base_font_size)
        if font_color:
            run.font.color.rgb = font_color

        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run.text = token[2:-2]
            run.bold = True
            run.italic = is_italic
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run.text = token[1:-1]
            run.italic = True
            run.bold = is_bold
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run.text = token[1:-1]
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 1)
            run.font.color.rgb = RGBColor(180, 40, 40)
        elif token.startswith('[') and ']' in token and '(' in token and token.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                label, url = m.groups()
                run.text = label
                run.font.color.rgb = RGBColor(30, 90, 180)
                run.underline = True
            else:
                run.text = token
        else:
            run.text = token
            run.bold = is_bold
            run.italic = is_italic

def markdown_to_docx(md_filepath, docx_filepath):
    with open(md_filepath, 'r', encoding='utf-8', errors='ignore') as f:
        content = f.read()

    doc = Document()

    # Set page margins (0.8 in)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    lines = content.split('\n')
    in_code_block = False
    code_block_lines = []
    in_table = False
    table_lines = []

    def flush_code_block():
        nonlocal code_block_lines, in_code_block
        if not code_block_lines:
            in_code_block = False
            return
        code_text = '\n'.join(code_block_lines)
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Inches(6.9)
        set_cell_background(cell, "F3F4F6")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(35, 45, 60)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        code_block_lines = []
        in_code_block = False

    def flush_table():
        nonlocal table_lines, in_table
        if not table_lines:
            in_table = False
            return
        
        parsed_rows = []
        for tl in table_lines:
            tl_strip = tl.strip()
            if not tl_strip:
                continue
            # Remove leading and trailing pipe
            if tl_strip.startswith('|'):
                tl_strip = tl_strip[1:]
            if tl_strip.endswith('|'):
                tl_strip = tl_strip[:-1]
            cols = [c.strip() for c in tl_strip.split('|')]
            # Check if this is separator row like |---|---|
            if all(re.match(r'^:?-+:?$', c) for c in cols if c):
                continue
            parsed_rows.append(cols)

        if parsed_rows:
            num_cols = max(len(r) for r in parsed_rows)
            # Pad rows
            for r in parsed_rows:
                while len(r) < num_cols:
                    r.append("")
            
            table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            
            for row_idx, row_data in enumerate(parsed_rows):
                is_header = (row_idx == 0)
                tr = table.rows[row_idx]
                for col_idx, cell_val in enumerate(row_data):
                    cell = tr.cells[col_idx]
                    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                    if is_header:
                        set_cell_background(cell, "1E3A8A") # Navy blue
                    else:
                        if row_idx % 2 == 1:
                            set_cell_background(cell, "F8FAFC")
                        else:
                            set_cell_background(cell, "FFFFFF")
                    
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    
                    clean_val = cell_val.replace('\\|', '|').strip()
                    if is_header:
                        run = p.add_run(clean_val)
                        run.bold = True
                        run.font.name = 'Calibri'
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        parse_inline_markdown(p, clean_val, base_font_size=9.5)
            
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        table_lines = []
        in_table = False

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                flush_code_block()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_block_lines = []
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Handle Markdown Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            flush_table()

        # Empty line
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^(\-{3,}|\*{3,}|_{3,})$', stripped):
            p = add_styled_paragraph(doc, space_before=6, space_after=6)
            run = p.add_run("_________________________________________________________________________________")
            run.font.color.rgb = RGBColor(200, 205, 215)
            run.font.size = Pt(9)
            i += 1
            continue

        # Heading 1
        if stripped.startswith('# '):
            h_text = stripped[2:].strip()
            p = add_styled_paragraph(doc, space_before=14, space_after=6)
            parse_inline_markdown(p, h_text, base_font_size=20, is_bold=True, font_color=RGBColor(24, 43, 73))
            i += 1
            continue

        # Heading 2
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = add_styled_paragraph(doc, space_before=12, space_after=4)
            parse_inline_markdown(p, h_text, base_font_size=15, is_bold=True, font_color=RGBColor(30, 64, 120))
            i += 1
            continue

        # Heading 3
        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = add_styled_paragraph(doc, space_before=10, space_after=3)
            parse_inline_markdown(p, h_text, base_font_size=12.5, is_bold=True, font_color=RGBColor(45, 85, 150))
            i += 1
            continue

        # Heading 4
        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = add_styled_paragraph(doc, space_before=8, space_after=2)
            parse_inline_markdown(p, h_text, base_font_size=11, is_bold=True, font_color=RGBColor(60, 100, 160))
            i += 1
            continue

        # Blockquote
        if stripped.startswith('> '):
            bq_text = stripped[2:].strip()
            p = add_styled_paragraph(doc, space_before=4, space_after=4)
            p.paragraph_format.left_indent = Inches(0.35)
            parse_inline_markdown(p, bq_text, base_font_size=10.5, is_italic=True, font_color=RGBColor(80, 90, 105))
            i += 1
            continue

        # Bullet lists
        bullet_match = re.match(r'^(\s*)([-*+])\s+(.*)$', line)
        if bullet_match:
            indent_spaces, _, item_text = bullet_match.groups()
            indent_level = len(indent_spaces) // 2
            p = add_styled_paragraph(doc, space_before=1, space_after=2)
            p.paragraph_format.left_indent = Inches(0.25 + 0.2 * indent_level)
            bullet_char = "• " if indent_level == 0 else "⁃ "
            run_b = p.add_run(bullet_char)
            run_b.font.name = 'Calibri'
            run_b.font.size = Pt(11)
            run_b.font.color.rgb = RGBColor(30, 64, 120)
            parse_inline_markdown(p, item_text, base_font_size=10.5)
            i += 1
            continue

        # Numbered lists
        num_match = re.match(r'^(\s*)(\d+[\.\)])\s+(.*)$', line)
        if num_match:
            indent_spaces, num_prefix, item_text = num_match.groups()
            indent_level = len(indent_spaces) // 2
            p = add_styled_paragraph(doc, space_before=1, space_after=2)
            p.paragraph_format.left_indent = Inches(0.25 + 0.2 * indent_level)
            run_n = p.add_run(num_prefix + " ")
            run_n.font.name = 'Calibri'
            run_n.font.size = Pt(10.5)
            run_n.bold = True
            run_n.font.color.rgb = RGBColor(30, 64, 120)
            parse_inline_markdown(p, item_text, base_font_size=10.5)
            i += 1
            continue

        # Regular paragraph
        p = add_styled_paragraph(doc, space_before=2, space_after=5)
        parse_inline_markdown(p, stripped, base_font_size=11)
        i += 1

    if in_code_block:
        flush_code_block()
    if in_table:
        flush_table()

    doc.save(docx_filepath)
    print(f"Successfully converted {md_filepath} -> {docx_filepath}")

if __name__ == "__main__":
    docs_dir = Path(r"c:\Users\Acer\Desktop\PolarGrid\docs")
    root_dir = Path(r"c:\Users\Acer\Desktop\PolarGrid")

    # Find all md files in docs
    md_files = list(docs_dir.glob("*.md"))
    # Also include root README.md
    root_readme = root_dir / "README.md"
    if root_readme.exists() and root_readme not in md_files:
        md_files.append(root_readme)

    print(f"Converting {len(md_files)} markdown files to docx format...")
    for mf in md_files:
        if mf.name == "README.md" and mf.parent == root_dir:
            out_name = "PROJECT_README.docx"
        else:
            out_name = mf.stem + ".docx"
        out_path = docs_dir / out_name
        try:
            markdown_to_docx(str(mf), str(out_path))
        except Exception as e:
            print(f"Error converting {mf}: {e}")

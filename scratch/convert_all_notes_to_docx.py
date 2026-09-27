import os
import re
import glob
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_styled_paragraph(doc, text, style=None, space_after=4, bold_prefix=None, is_quote=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    if is_quote:
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.bold = True
        r_pre.font.size = Pt(11)
        r_pre.font.color.rgb = RGBColor(14, 75, 133)

    # Process bold markup **text** and code markup `code`
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            inner = token[2:-2]
            run = p.add_run(inner)
            run.font.name = "Calibri"
            run.font.bold = True
            run.font.size = Pt(10.5)
            if is_quote:
                run.font.italic = True
        elif token.startswith("`") and token.endswith("`"):
            inner = token[1:-1]
            run = p.add_run(inner)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 40, 40)
        else:
            run = p.add_run(token)
            run.font.name = "Calibri"
            run.font.size = Pt(10.5)
            if is_quote:
                run.font.italic = True
                run.font.color.rgb = RGBColor(70, 70, 70)

    return p

def convert_markdown_file_to_docx(md_path, docx_path, document_title=None):
    doc = docx.Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Polar Grid — SIH Notes | Mawson Station, Antarctica")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(140, 140, 140)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Confidential — Smart India Hackathon (SIH) Official Team Notes")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(140, 140, 140)

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code_block = False
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        num_cols = max(len(r) for r in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        for row_idx, row_data in enumerate(table_rows):
            is_header = (row_idx == 0)
            row = table.rows[row_idx]
            for col_idx in range(num_cols):
                cell = row.cells[col_idx]
                cell_text = row_data[col_idx] if col_idx < len(row_data) else ""
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                
                if is_header:
                    set_cell_background(cell, "0E4B85") # Deep Blue
                elif row_idx % 2 == 1:
                    set_cell_background(cell, "F4F8FB") # Soft Ice Blue
                else:
                    set_cell_background(cell, "FFFFFF")

                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                
                # Parse formatting inside cell
                tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', cell_text)
                for token in tokens:
                    if not token:
                        continue
                    if token.startswith("**") and token.endswith("**"):
                        r = p.add_run(token[2:-2])
                        r.font.bold = True
                    elif token.startswith("`") and token.endswith("`"):
                        r = p.add_run(token[1:-1])
                        r.font.name = "Consolas"
                        r.font.size = Pt(9)
                    else:
                        r = p.add_run(token)

                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5) if not is_header else Pt(10)
                    if is_header:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        r.font.color.rgb = RGBColor(30, 30, 30)

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(6)
        table_rows = []
        in_table = False

    i = 0
    while i < len(lines):
        raw_line = lines[i]
        line = raw_line.rstrip()

        # Check code fence
        if line.startswith("```"):
            if in_table:
                flush_table()
            in_code_block = not in_code_block
            i += 1
            continue

        if in_code_block:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            r = p.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(40, 40, 40)
            i += 1
            continue

        # Check markdown table line
        if line.startswith("|") and line.endswith("|"):
            # Check if it's separator row | :--- | :--- |
            if re.match(r'^\|[\s:-]+\|', line):
                i += 1
                continue
            cells = [c.strip() for c in line.split("|")[1:-1]]
            if not in_table:
                in_table = True
                table_rows = []
            table_rows.append(cells)
            i += 1
            continue
        elif in_table:
            flush_table()

        # Empty line
        if not line:
            i += 1
            continue

        # Horizontal rule ---
        if line.startswith("---") or line.startswith("==="):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run("―" * 45)
            r.font.name = "Calibri"
            r.font.color.rgb = RGBColor(200, 210, 220)
            i += 1
            continue

        # Heading 1 (# ...)
        if line.startswith("# "):
            text = line[2:].strip()
            h = doc.add_heading(text, level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            for r in h.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(18)
                r.font.bold = True
                r.font.color.rgb = RGBColor(14, 75, 133) # Deep Blue
            i += 1
            continue

        # Heading 2 (## ...)
        if line.startswith("## "):
            text = line[3:].strip()
            h = doc.add_heading(text, level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(3)
            for r in h.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(14)
                r.font.bold = True
                r.font.color.rgb = RGBColor(2, 132, 199) # Sky Blue
            i += 1
            continue

        # Heading 3 (### ...)
        if line.startswith("### "):
            text = line[4:].strip()
            h = doc.add_heading(text, level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            for r in h.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(11.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(40, 60, 80)
            i += 1
            continue

        # Blockquote (> ...)
        if line.startswith("> "):
            text = line[2:].strip()
            # Accumulate multi-line quotes
            while i + 1 < len(lines) and lines[i+1].startswith("> "):
                i += 1
                text += " " + lines[i][2:].strip()
            add_styled_paragraph(doc, text, is_quote=True)
            i += 1
            continue

        # Bullet list item (- ... or * ...)
        if line.startswith("- ") or line.startswith("* "):
            text = line[2:].strip()
            add_styled_paragraph(doc, text, style="List Bullet", space_after=2)
            i += 1
            continue

        # Numbered list item (1. ... or 2. ...)
        num_match = re.match(r'^(\d+)\.\s+(.*)', line)
        if num_match:
            prefix = num_match.group(1) + ". "
            text = num_match.group(2)
            add_styled_paragraph(doc, prefix + text, style="List Number", space_after=2)
            i += 1
            continue

        # Standard paragraph
        add_styled_paragraph(doc, line, space_after=4)
        i += 1

    if in_table:
        flush_table()

    doc.save(docx_path)
    file_size = os.path.getsize(docx_path)
    print(f"Generated {docx_path} ({file_size} bytes)")

def convert_all():
    notes_dir = os.path.abspath(os.path.join("docs", "sih_notes"))
    md_files = sorted(glob.glob(os.path.join(notes_dir, "*.md")))
    print(f"Found {len(md_files)} module notes in {notes_dir}")

    for md_file in md_files:
        docx_file = md_file.replace(".md", ".docx")
        convert_markdown_file_to_docx(md_file, docx_file)

    # Also convert master combined document
    master_md = os.path.abspath(os.path.join("docs", "POLAR_GRID_FINAL_SIH_NOTES.md"))
    master_docx = os.path.abspath(os.path.join("docs", "POLAR_GRID_FINAL_SIH_NOTES.docx"))
    if os.path.exists(master_md):
        print(f"Converting master document: {master_md}")
        convert_markdown_file_to_docx(master_md, master_docx)

if __name__ == "__main__":
    convert_all()

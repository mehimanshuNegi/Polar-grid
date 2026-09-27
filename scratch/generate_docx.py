import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def markdown_to_docx(md_path, docx_path):
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code_block = False
    
    for raw_line in lines:
        line = raw_line.rstrip()
        
        # Check code block fences
        if line.startswith("```"):
            in_code_block = not in_code_block
            continue
            
        if in_code_block:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(60, 60, 60)
            continue
            
        # Empty line
        if not line:
            continue
            
        # Header 1 (# ...)
        if line.startswith("# "):
            text = line[2:].strip()
            h = doc.add_heading(text, level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            for r in h.runs:
                r.font.name = "Arial"
                r.font.color.rgb = RGBColor(14, 75, 133)
            continue

        # Header 2 (## ...)
        if line.startswith("## "):
            text = line[3:].strip()
            h = doc.add_heading(text, level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            for r in h.runs:
                r.font.name = "Arial"
                r.font.color.rgb = RGBColor(24, 115, 186)
            continue

        # Header 3 (### ...)
        if line.startswith("### "):
            text = line[4:].strip()
            h = doc.add_heading(text, level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(2)
            for r in h.runs:
                r.font.name = "Arial"
                r.font.color.rgb = RGBColor(50, 50, 50)
            continue

        # Blockquote (> ...)
        if line.startswith("> "):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.italic = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(80, 80, 80)
            continue

        # Bullet list item (- ... or * ...)
        if line.startswith("- ") or line.startswith("* "):
            text = line[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            # Parse simple bolding
            parts = text.split("**")
            for i, part in enumerate(parts):
                run = p.add_run(part)
                run.font.name = "Arial"
                run.font.size = Pt(10)
                if i % 2 == 1:
                    run.font.bold = True
            continue

        # Numbered list item (1. ... or 2. ...)
        if len(line) > 2 and line[0].isdigit() and line[1:3] in [". ", ": "]:
            text = line[3:].strip()
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            parts = text.split("**")
            for i, part in enumerate(parts):
                run = p.add_run(part)
                run.font.name = "Arial"
                run.font.size = Pt(10)
                if i % 2 == 1:
                    run.font.bold = True
            continue

        # Standard paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        parts = line.split("**")
        for i, part in enumerate(parts):
            run = p.add_run(part)
            run.font.name = "Arial"
            run.font.size = Pt(10.5)
            if i % 2 == 1:
                run.font.bold = True

    doc.save(docx_path)
    print(f"Successfully generated DOCX at {docx_path} ({os.path.getsize(docx_path)} bytes)")

if __name__ == "__main__":
    md_file = os.path.abspath(os.path.join("docs", "POLAR_GRID_FINAL_SIH_NOTES.md"))
    docx_file = os.path.abspath(os.path.join("docs", "POLAR_GRID_FINAL_SIH_NOTES.docx"))
    markdown_to_docx(md_file, docx_file)
